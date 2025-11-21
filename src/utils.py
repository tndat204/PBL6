from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import tempfile

def export_results_to_excel(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Match Results"

    # Header
    headers = [
        "CV Name/URL",
        "TechnicalSkills",
        "SoftSkills",
        "Experience",
        "Education",
        "Other",
        "TotalScore",
        "Raw Technical Skills",
        "Raw Soft Skills",
        "Raw Experience",
        "Raw Education",
        "Raw Certifications",
        "Raw Languages",
        "Raw OtherInfo"
    ]
    ws.append(headers)

    # Rows
    for r in results:
        cv_label = r.get("cv_filename") or r.get("cv_url") or "Unknown"

        # Nếu CV lỗi → ghi 1 dòng
        if "error" in r:
            ws.append([cv_label, "ERROR", r["error"]])
            continue

        score = r["match_score"]
        data = r["cv_data"]

        row = [
            cv_label,
            score.get("TechnicalSkills"),
            score.get("SoftSkills"),
            score.get("Experience"),
            score.get("Education"),
            score.get("Other"),
            score.get("TotalScore"),
            str(data.get("TechnicalSkills")),
            str(data.get("SoftSkills")),
            str(data.get("Experience")),
            str(data.get("Education")),
            str(data.get("Certifications")),
            str(data.get("Languages")),
            str(data.get("OtherInfo"))
        ]

        ws.append(row)

    # Auto width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2

    # Save temp file
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    return tmp.name
